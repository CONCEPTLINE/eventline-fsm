"""Text aus Dateien holen — PDF, DOCX, XLSX, TXT/CSV, Bilder — alles im Haus.

PDF-Reihenfolge: eingebetteter Text (pypdf, dann pdfplumber) → unter 50 Zeichen
pro Seite gilt die Datei als Scan → OCR mit ocrmypdf (--skip-text, Tesseract
deu+eng), Notnagel pdftoppm + Tesseract. Jede Zwischendatei entsteht unter
TMPDIR, also im Tresor, und wird am Ende gelöscht.
"""

from __future__ import annotations

import base64
import io
import logging
import re
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path

from .config import Einstellungen

log = logging.getLogger("ki.extract")

# libmagic einmal laden; fehlt es (Entwicklungsrechner), entscheiden Endung und Angabe.
try:
    import magic as _magic  # python-magic, libmagic1 im Image
except Exception:  # noqa: BLE001
    _magic = None

MIN_ZEICHEN_PRO_SEITE = 50
OCR_SPRACHEN = "deu+eng"

MIME_PDF = "application/pdf"
MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
BILD_MIMES = {"image/jpeg", "image/png", "image/tiff", "image/webp", "image/bmp", "image/gif"}
TEXT_MIMES = {"text/plain", "text/csv", "text/markdown", "text/html", "application/json", "application/xml", "text/xml"}

ENDUNG_MIME = {
    ".pdf": MIME_PDF,
    ".docx": MIME_DOCX,
    ".xlsx": MIME_XLSX,
    ".txt": "text/plain",
    ".csv": "text/csv",
    ".md": "text/markdown",
    ".json": "application/json",
    ".xml": "application/xml",
    ".log": "text/plain",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
    ".gif": "image/gif",
    ".heic": "image/heic",
}


class ExtraktFehler(Exception):
    """Klartext, warum eine Datei nicht gelesen werden konnte."""


@dataclass
class Extrakt:
    text: str
    mime: str
    seiten: int | None = None
    ocr: bool = False
    gekuerzt: bool = False
    hinweis: str | None = None
    ist_bild: bool = False


# --- Erkennung ----------------------------------------------------------------


def erkenne_mime(pfad: Path, angabe: str | None, dateiname: str) -> str:
    """Inhalt (libmagic) schlägt Endung schlägt Angabe — Angaben aus dem Browser sind oft generisch."""
    endung = ""
    m = re.search(r"\.[A-Za-z0-9]{1,8}$", dateiname or pfad.name)
    if m:
        endung = m.group(0).lower()
    aus_endung = ENDUNG_MIME.get(endung)
    erkannt = ""
    if _magic is not None:
        try:
            erkannt = str(_magic.from_file(str(pfad), mime=True) or "").lower()
        except Exception:  # noqa: BLE001 — libmagic streikt: Endung/Angabe reichen
            erkannt = ""
    if erkannt and erkannt not in ("application/octet-stream", "text/plain", "application/zip"):
        return erkannt
    if aus_endung:
        return aus_endung
    if erkannt:
        return erkannt
    if angabe and "/" in angabe:
        return angabe.lower()
    return "application/octet-stream"


# --- Haupteinstieg ------------------------------------------------------------


def extrahiere(pfad: Path, mime_angabe: str | None, dateiname: str, cfg: Einstellungen) -> Extrakt:
    groesse = pfad.stat().st_size
    if groesse > cfg.max_datei_mb * 1024 * 1024:
        raise ExtraktFehler(f"Datei ist grösser als {cfg.max_datei_mb} MB")
    if groesse == 0:
        raise ExtraktFehler("Datei ist leer")
    mime = erkenne_mime(pfad, mime_angabe, dateiname)

    seiten: int | None = None
    ocr = False
    ist_bild = False
    hinweis: str | None = None

    if mime == MIME_PDF:
        text, seiten, ocr = _pdf(pfad, cfg)
    elif mime == MIME_DOCX:
        text = _docx(pfad)
    elif mime == MIME_XLSX:
        text = _xlsx(pfad)
    elif mime in TEXT_MIMES or mime.startswith("text/"):
        text = _text(pfad, cfg.max_text_zeichen)
    elif mime in BILD_MIMES:
        text = _bild_ocr(pfad, cfg)
        seiten, ocr, ist_bild = 1, True, True
    elif mime == "image/heic":
        raise ExtraktFehler("HEIC-Bilder können nicht gelesen werden — bitte als JPG oder PNG ablegen")
    else:
        text = ""
        hinweis = f"Dateityp {mime} kann nicht gelesen werden — nur der Dateiname ist bekannt"

    text = _normalisiere(text)
    text, gekuerzt = _kappen(text, cfg.max_text_zeichen)
    if not text.strip() and hinweis is None:
        hinweis = "Kein lesbarer Text gefunden (leeres Dokument oder OCR ohne Ergebnis)"
    return Extrakt(text=text, mime=mime, seiten=seiten, ocr=ocr, gekuerzt=gekuerzt, hinweis=hinweis, ist_bild=ist_bild)


def _normalisiere(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    zeilen = [re.sub(r"[ \t]+", " ", z).strip() for z in text.split("\n")]
    text = "\n".join(zeilen)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _kappen(text: str, max_zeichen: int) -> tuple[str, bool]:
    if len(text) <= max_zeichen:
        return text, False
    # Tausender mit Apostroph (Schweizer Schreibweise), nur in den Zahlen — nie im Text.
    von = f"{max_zeichen:,}".replace(",", "'")
    gesamt = f"{len(text):,}".replace(",", "'")
    return f"{text[:max_zeichen]}\n\n[Text gekürzt: nur die ersten {von} von {gesamt} Zeichen]", True


# --- PDF ----------------------------------------------------------------------


def _pdf(pfad: Path, cfg: Einstellungen) -> tuple[str, int | None, bool]:
    texte, seiten = _pdf_pypdf(pfad, cfg.pdf_max_seiten)
    ocr = False
    if _zu_wenig(texte):
        alt = _pdf_pdfplumber(pfad, min(seiten or 60, 60))
        if alt and not _zu_wenig(alt):
            texte = alt
    if _zu_wenig(texte):
        ocr_texte = _pdf_ocr(pfad, cfg)
        if ocr_texte:
            texte = ocr_texte
            ocr = True
    return _seiten_zusammen(texte), seiten, ocr


def _seiten_zusammen(texte: list[str]) -> str:
    teile = []
    for i, t in enumerate(texte, start=1):
        if t.strip():
            teile.append(f"[Seite {i}]\n{t.strip()}")
    return "\n\n".join(teile)


def _zu_wenig(texte: list[str]) -> bool:
    if not texte:
        return True
    zeichen = sum(len(t.strip()) for t in texte)
    return zeichen / len(texte) < MIN_ZEICHEN_PRO_SEITE


def _pdf_pypdf(pfad: Path, max_seiten: int) -> tuple[list[str], int | None]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(pfad))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ExtraktFehler("PDF ist passwortgeschützt")
            except ExtraktFehler:
                raise
            except Exception as e:  # noqa: BLE001
                raise ExtraktFehler("PDF ist passwortgeschützt") from e
        seiten = len(reader.pages)
        texte = []
        for i, seite in enumerate(reader.pages):
            if i >= max_seiten:
                break
            try:
                texte.append(seite.extract_text() or "")
            except Exception:  # noqa: BLE001 — einzelne kaputte Seite bricht nicht alles ab
                texte.append("")
        return texte, seiten
    except ExtraktFehler:
        raise
    except Exception as e:  # noqa: BLE001
        log.warning("pypdf konnte die Datei nicht lesen (%s)", type(e).__name__)
        return [], None


def _pdf_pdfplumber(pfad: Path, max_seiten: int) -> list[str]:
    try:
        import pdfplumber

        texte = []
        with pdfplumber.open(str(pfad)) as pdf:
            for i, seite in enumerate(pdf.pages):
                if i >= max_seiten:
                    break
                try:
                    texte.append(seite.extract_text() or "")
                except Exception:  # noqa: BLE001
                    texte.append("")
        return texte
    except Exception as e:  # noqa: BLE001
        log.warning("pdfplumber konnte die Datei nicht lesen (%s)", type(e).__name__)
        return []


def _pdf_ocr(pfad: Path, cfg: Einstellungen) -> list[str]:
    n = max(1, cfg.ocr_max_seiten)
    cfg.tmp_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=cfg.tmp_dir, prefix="ocr-") as d:
        ziel = Path(d) / "ocr.pdf"
        cmd = [
            "ocrmypdf", "--skip-text", "-l", OCR_SPRACHEN, "--output-type", "pdf",
            "--optimize", "0", "--jobs", "2", "--pages", f"1-{n}", "--quiet",
            str(pfad), str(ziel),
        ]
        try:
            p = subprocess.run(cmd, capture_output=True, text=True, timeout=cfg.ocr_timeout_s)
            code: int | str = p.returncode
        except (OSError, subprocess.SubprocessError) as e:
            code = type(e).__name__
        if code == 0 and ziel.exists():
            texte, _ = _pdf_pypdf(ziel, n)
            if not _zu_wenig(texte):
                return texte
            alt = _pdf_pdfplumber(ziel, n)
            if not _zu_wenig(alt):
                return alt
        else:
            log.warning("ocrmypdf fehlgeschlagen (%s) — Notnagel pdftoppm + Tesseract", code)
        return _pdf_rastern_ocr(pfad, Path(d), n, cfg)


def _pdf_rastern_ocr(pfad: Path, ordner: Path, n: int, cfg: Einstellungen) -> list[str]:
    praefix = ordner / "seite"
    cmd = ["pdftoppm", "-r", "200", "-gray", "-png", "-f", "1", "-l", str(n), str(pfad), str(praefix)]
    try:
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=cfg.ocr_timeout_s)
    except (OSError, subprocess.SubprocessError) as e:
        log.warning("pdftoppm fehlgeschlagen (%s)", type(e).__name__)
        return []
    if p.returncode != 0:
        log.warning("pdftoppm fehlgeschlagen (Code %s)", p.returncode)
        return []
    return [_tesseract(b, cfg) for b in sorted(ordner.glob("seite-*.png"))]


# --- Bilder -------------------------------------------------------------------


def _bild_oeffnen(pfad: Path):
    from PIL import Image, ImageOps, UnidentifiedImageError

    try:
        img = Image.open(pfad)
        img.load()
    except (UnidentifiedImageError, OSError) as e:
        raise ExtraktFehler("Bild kann nicht gelesen werden (Format nicht unterstützt oder Datei beschädigt)") from e
    return ImageOps.exif_transpose(img) or img


def _tesseract(bild_pfad: Path, cfg: Einstellungen) -> str:
    try:
        import pytesseract

        img = _bild_oeffnen(bild_pfad).convert("L")
        if max(img.size) > 3000:
            img.thumbnail((3000, 3000))
        return pytesseract.image_to_string(img, lang=OCR_SPRACHEN, timeout=cfg.ocr_timeout_s) or ""
    except ExtraktFehler:
        raise
    except Exception as e:  # noqa: BLE001 — Tesseract-Zeitüberschreitung o. ä.: Seite bleibt leer
        log.warning("Tesseract fehlgeschlagen (%s)", type(e).__name__)
        return ""


def _bild_ocr(pfad: Path, cfg: Einstellungen) -> str:
    _bild_oeffnen(pfad)  # früh scheitern, wenn das Bild unlesbar ist
    return _tesseract(pfad, cfg)


def bild_base64(pfad: Path, max_px: int) -> str:
    """JPEG-Verkleinerung fürs Bildmodell (kleiner = schneller, Lesbarkeit reicht)."""
    img = _bild_oeffnen(pfad).convert("RGB")
    img.thumbnail((max_px, max_px))
    puffer = io.BytesIO()
    img.save(puffer, format="JPEG", quality=85, optimize=True)
    return base64.b64encode(puffer.getvalue()).decode("ascii")


# --- Office / Text ------------------------------------------------------------


def _docx(pfad: Path) -> str:
    try:
        import docx
    except ImportError as e:  # pragma: no cover
        raise ExtraktFehler("DOCX-Unterstützung fehlt") from e
    try:
        dokument = docx.Document(str(pfad))
    except Exception as e:  # noqa: BLE001
        raise ExtraktFehler("DOCX kann nicht gelesen werden") from e
    teile: list[str] = []
    # Kopf-/Fusszeilen zuerst: dort stehen Briefkopf, Partei und Nummern.
    for abschnitt in dokument.sections:
        for bereich in (abschnitt.header, abschnitt.footer):
            try:
                for p in bereich.paragraphs:
                    if p.text.strip():
                        teile.append(p.text.strip())
            except Exception:  # noqa: BLE001
                pass
    for p in dokument.paragraphs:
        if p.text.strip():
            teile.append(p.text.strip())
    for tabelle in dokument.tables:
        for zeile in tabelle.rows:
            zellen = []
            for zelle in zeile.cells:
                t = zelle.text.strip()
                if t and (not zellen or zellen[-1] != t):
                    zellen.append(t)
            if zellen:
                teile.append(" | ".join(zellen))
    return "\n".join(teile)


def _xlsx(pfad: Path, max_zellen: int = 20_000) -> str:
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover
        raise ExtraktFehler("XLSX-Unterstützung fehlt") from e
    try:
        wb = openpyxl.load_workbook(str(pfad), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise ExtraktFehler("XLSX kann nicht gelesen werden") from e
    teile: list[str] = []
    zellen = 0
    try:
        for ws in wb.worksheets:
            teile.append(f"[Blatt {ws.title}]")
            for zeile in ws.iter_rows(values_only=True):
                werte = [str(v).strip() for v in zeile if v is not None and str(v).strip()]
                zellen += len(werte)
                if werte:
                    teile.append(" | ".join(werte))
                if zellen >= max_zellen:
                    teile.append("[weitere Zellen ausgelassen]")
                    break
            if zellen >= max_zellen:
                break
    finally:
        wb.close()
    return "\n".join(teile)


def _text(pfad: Path, max_zeichen: int) -> str:
    roh = pfad.read_bytes()[: max_zeichen * 4]
    for codierung in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return roh.decode(codierung)
        except UnicodeDecodeError:
            continue
    return roh.decode("utf-8", errors="replace")
